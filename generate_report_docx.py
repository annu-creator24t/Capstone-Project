"""Script to generate a comprehensive, publication-quality Microsoft Word (.docx) report
for Capstone Project Phase III & Mid-Term Progress Details.
Fulfills all 9 evaluation requirements specified for college submission.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell internal margins (padding) in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def style_table_header(row, col_widths=None):
    """Apply styling to the header row of a table."""
    for i, cell in enumerate(row.cells):
        set_cell_background(cell, "1A365D")  # Deep Navy
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9.5)
                r.font.name = "Calibri"
        if col_widths and i < len(col_widths):
            cell.width = col_widths[i]

def style_table_row(row, is_even=False, col_widths=None):
    """Apply styling to data rows of a table."""
    bg_color = "F7FAFC" if is_even else "FFFFFF"
    for i, cell in enumerate(row.cells):
        if bg_color != "FFFFFF":
            set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)
                r.font.name = "Calibri"
                r.font.color.rgb = RGBColor(45, 55, 72)
        if col_widths and i < len(col_widths):
            cell.width = col_widths[i]

def add_callout(doc, text, title=None, border_color="3182CE", bg_color="EBF8FF"):
    """Add a styled callout box with a colored left accent border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(10)
        run_title.font.color.rgb = RGBColor(26, 54, 93)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(45, 55, 72)
    
    # Add small spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def build_word_report(output_path):
    print(f"Generating Comprehensive Word Report at: {output_path}...")
    doc = Document()
    
    # Page setup: Standard Letter, 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Configure Header & Footer
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("CAPSTONE PROJECT PHASE III & MID-TERM PROGRESS REPORT")
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(8)
        r_hdr.font.color.rgb = RGBColor(113, 128, 150)
        
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ftr = p_ftr.add_run("Cloud-Based Digital Twin for Connected Vehicle Diagnostics | Confidential & Academic Use")
        r_ftr.font.name = "Calibri"
        r_ftr.font.size = Pt(8)
        r_ftr.font.color.rgb = RGBColor(160, 174, 192)

    # -------------------------------------------------------------------------
    # COVER / HEADER BLOCK
    # -------------------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("CAPSTONE PROJECT PHASE III &\nMID-TERM PROGRESS REPORT")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(26, 54, 93)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Cloud-Based Digital Twin for Connected Vehicle Diagnostics & Performance Analytics")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.bold = True
    run_sub.font.color.rgb = RGBColor(43, 108, 176)

    # Metadata Box
    meta_text = (
        "• Academic Level: B.Tech / Final-Year Capstone Project\n"
        "• Evaluation Milestone: Phase III & Mid-Term Progress Review\n"
        "• Proposed Research Direction: Uncertainty- and Cost-Aware Active Digital Twin for Connected Vehicle Diagnostics Under Intermittent Connectivity\n"
        "• Git Repository Status: Synced on branch main (Remote Commit: 46996fa)\n"
        "• Verified Test Suite: 171 of 171 Tests Passing (100% Pass Rate in 1.47s)\n"
        "• Date: October 2026 | Document Status: Verified & College Submission-Ready"
    )
    add_callout(doc, meta_text, title="PROJECT IDENTIFICATION & VERIFICATION SUMMARY", border_color="1A365D", bg_color="F7FAFC")

    # -------------------------------------------------------------------------
    # EXECUTIVE SUMMARY
    # -------------------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r = h1.add_run("Executive Summary")
    r.font.color.rgb = RGBColor(26, 54, 93)
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "Modern connected and autonomous vehicles (CAVs) operate in dynamic, non-stationary wireless environments characterized by variable transmission latency (10 to 600 ms), packet loss (up to 50%), and telemetry bandwidth constraints. When powertrain anomalies arise—such as rising engine or coolant temperatures—conventional onboard diagnostic systems (OBD-II) rely on rigid, static threshold alerts that produce unacceptable false alarm rates during transient engine loading and cannot isolate root causes. Meanwhile, existing cloud-based passive digital twin architectures continuously ingest low-frequency telemetry but struggle to distinguish between physically distinct faults that produce identical passive symptoms (for example, a cooling fan actuator relay failure versus a temperature sensor calibration bias versus a coolant fluid leak).\n\n"
        "This capstone project designs, implements, and empirically validates an Active Digital Twin (ADT) diagnostic framework. The system combines: (1) a physical vehicle powertrain and thermal dynamic simulator governed by coupled first-order differential equations and nonlinear thermostat hysteresis; (2) a non-stationary wireless network emulator modeling delay, jitter, Bernoulli packet loss, priority queueing, bandwidth serialization, and Age-of-Information (AoI) freshness; (3) a cloud-based nominal physics twin that reconstructs expected vehicle states, computes continuous and categorical residuals, filters telemetry by freshness qualification, and evaluates anomalies using two-level thresholds with persistence counters; and (4) an interactive research analytics dashboard (Streamlit on port 8501) and high-speed telemetry cockpit (HTML5/WebSocket on port 8050).\n\n"
        "As of Phase III, all 171 automated unit and integration tests execute successfully with zero failures (100% pass rate in 1.47s). An extensive 93-run research campaign across eight distinct impairment sweeps confirms a 91.80% empirical packet delivery ratio, a mean Age of Information of 0.1672 seconds, a 93.3% true positive fault detection rate, and a 0.0% false alarm rate across healthy baseline driving. Furthermore, a 40-second controlled experiment verifies immediate detection of cooling fan actuator failure (t = 0.50s) with a +5.65°C thermal divergence."
    )

    # -------------------------------------------------------------------------
    # REQUIREMENT 1: SYSTEM ARCHITECTURE AND DESIGN
    # -------------------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. System Architecture and Design")
    r.font.color.rgb = RGBColor(26, 54, 93)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "The project implements an Edge-Cloud Cyber-Physical Twin Architecture organized into four decoupled tiers. The system separates the physical vehicle domain from the communication channel and cloud diagnostic reasoning engine."
    )

    doc.add_heading("1.1 Architecture Decomposition & Module Responsibilities", level=2)
    
    # Table of Modules
    tbl = doc.add_table(rows=5, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.2), Inches(1.5), Inches(1.2), Inches(2.6)]
    
    headers = ["Tier / Domain", "Module Directory", "Primary Classes", "System Responsibilities"]
    for i, h in enumerate(headers):
        tbl.cell(0, i).paragraphs[0].text = h
    style_table_header(tbl.rows[0], widths)
    
    modules_data = [
        ("Tier 1: Edge Vehicle Simulator", "simulator/", "VehicleModel, VehicleState, FaultInjector, SensorGenerator", "Simulates vehicle velocity kinematics, dynamic engine load, and continuous thermal energy balance dQ_in - dQ_out with thermostat fan hysteresis. Injects controlled synthetic faults and generates realistic Gaussian sensor noise."),
        ("Tier 2: Network Impairment Layer", "network/", "NetworkEmulator, DelayModel, PacketLossModel, PacketPriorityQueue, AoITracker", "Emulates cellular network impairments between vehicle and cloud. Applies stochastic delays, Bernoulli packet drops, priority queueing, and bandwidth limits. Tracks real-time Age-of-Information (AoI) per sensor stream."),
        ("Tier 3: Cloud Digital Twin", "digital_twin/", "NominalVehicleModel, StateEstimator, ResidualGenerator, TelemetryFreshness, AnomalyDetector", "Reconstructs nominal vehicle behavior in the cloud, buffers delayed telemetry, computes continuous and categorical residuals r(t) = y_obs(t) - y_hat(t), qualifies freshness states, and executes persistence-gated anomaly detection."),
        ("Tier 4: Presentation & Analytics", "dashboard/ & app.py", "DashboardSimulationSession, LiveSimulationEngine, Streamlit App", "Provides dual interfaces: (1) Streamlit analytics cockpit on port 8501 with real-time gauges, residual plots, network tuning, and campaign explorer; (2) HTML5 cockpit on port 8050 with responsive gauges and telemetry streaming.")
    ]
    
    for row_idx, data in enumerate(modules_data, start=1):
        row = tbl.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
        style_table_row(row, is_even=(row_idx % 2 == 0), col_widths=widths)

    doc.add_heading("1.2 Full-Stack Technology Manifest", level=2)
    
    tbl_tech = doc.add_table(rows=8, cols=4)
    tbl_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths_tech = [Inches(1.4), Inches(1.5), Inches(1.1), Inches(2.5)]
    
    headers_tech = ["Category", "Technology / Framework", "Version", "Technical Purpose"]
    for i, h in enumerate(headers_tech):
        tbl_tech.cell(0, i).paragraphs[0].text = h
    style_table_header(tbl_tech.rows[0], widths_tech)
    
    tech_data = [
        ("Runtime", "Python (64-bit)", "3.10.0", "Core execution environment for all simulator, network, and twin logic"),
        ("Validation", "Pydantic", "2.13.4", "Strict type-safe configuration schemas for faults, network, and twin"),
        ("Numerics", "NumPy & SciPy", "2.2.6 / 1.15.3", "Numerical differential equation updates and statistical distributions"),
        ("Data Handling", "Pandas", "2.3.3", "Telemetry history buffering, time-series dataframes, and CSV parsing"),
        ("Testing Harness", "Pytest & Unittest", "8.4.1", "Automated execution of 171 unit, integration, and regression test cases"),
        ("Research Dashboard", "Streamlit & Plotly", "1.54.0 / 6.5.2", "Interactive engineering analytics UI, dynamic gauges, and sweep charts"),
        ("Report Generation", "ReportLab & python-docx", "4.4.3 / 1.2.0", "Automated publication-quality PDF guide and Word submission reporting")
    ]
    for row_idx, data in enumerate(tech_data, start=1):
        row = tbl_tech.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
        style_table_row(row, is_even=(row_idx % 2 == 0), col_widths=widths_tech)

    # -------------------------------------------------------------------------
    # REQUIREMENT 2: METHODOLOGY AND IMPLEMENTATION
    # -------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("2. Methodology and Implementation")
    r.font.color.rgb = RGBColor(26, 54, 93)

    doc.add_heading("2.1 Mathematical Formulation of Vehicle Thermal Powertrain Dynamics", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "The engine thermal model implemented in simulator/vehicle_model.py is governed by continuous-time energy conservation evaluated over discrete simulation intervals dt = 0.1s:\n\n"
        "    dT_eng / dt = (Q_in - Q_out) / C_thermal\n\n"
        "Where C_thermal = 4500.0 J/°C is the lumped engine block and coolant heat capacity.\n"
        "1. Dynamic Heat Generation Rate (Q_in):\n"
        "    Q_in = (k_base + k_load * L) * (RPM / 1000.0)\n"
        "    with k_base = 180.0 W, k_load = 650.0 W, and engine load L in [0.0, 1.0].\n"
        "2. Multi-Factor Heat Dissipation Rate (Q_out):\n"
        "    Q_out = (k_nat + k_speed * v + k_fan * F) * (T_eng - T_ambient)\n"
        "    with natural convection k_nat = 12.0 W/°C, ram-air cooling k_speed = 0.8 W/(km/h·°C), fan cooling k_fan = 45.0 W/°C, ambient temperature T_ambient = 25.0°C, and vehicle speed v in km/h.\n"
        "3. ECU Thermostat Cooling Fan Relay Hysteresis:\n"
        "    F(t) = 1 (Fan ON) if T_eng >= 95.0°C\n"
        "    F(t) = 0 (Fan OFF) if T_eng <= 90.0°C\n"
        "    F(t) = F(t - dt) (Hold State) if 90.0°C < T_eng < 95.0°C"
    )

    doc.add_heading("2.2 Age of Information (AoI) Continuous Sawtooth Integration", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "Unlike conventional network latency (which only measures the transmission delay of received packets), Age of Information (AoI) quantifies data staleness from the receiver perspective. Let u_s(t) denote the generation timestamp of the most recent packet received for sensor s by time t. Instantaneous AoI is defined as:\n\n"
        "    AoI_s(t) = t - u_s(t)\n\n"
        "When no packet arrives, AoI grows linearly with unit slope (d AoI / dt = 1). When a new packet arrives at time t_arr with generation timestamp t_gen > u_s(t), the AoI drops non-linearly to the packet transfer delay (t_arr - t_gen). The network/aoi_tracker.py module implements continuous trapezoidal integration of the area under the AoI curve to compute the continuous time-average AoI:\n\n"
        "    Average_AoI_s = (1 / T) * integral_0^T AoI_s(t) dt = (1 / T) * sum_i=1^M [0.5 * (AoI_s(t_start) + AoI_s(t_end)) * dt_i]"
    )

    doc.add_heading("2.3 Cloud Digital Twin Residual Generation & Persistence Filtering", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "Implemented across digital_twin/residual_generator.py and digital_twin/anomaly_detector.py:\n"
        "• Raw Discrepancy: r_temp(t) = T_obs(t) - T_hat(t)\n"
        "• Defensive Relative Residual: r_rel(t) = |r_temp(t)| / max(|T_hat(t)|, epsilon), with epsilon = 1e-4 preventing division-by-zero.\n"
        "• Categorical Actuator Mismatch: F_obs(t) != F_hat(t) flags an immediate cooling fan actuator relay failure when the ECU thermostat expects the fan to be running.\n"
        "• Two-Level Thresholding: Warning threshold at |r_temp| >= 2.0°C; Anomaly threshold at |r_temp| >= 5.0°C.\n"
        "• Persistence Gating Algorithm: An internal counter C increments for consecutive threshold exceedances. A confirmed anomaly is declared if and only if C >= N_persist = 3 consecutive evaluation cycles, effectively filtering transient noise and single packet spikes."
    )

    # -------------------------------------------------------------------------
    # REQUIREMENT 3: PROTOTYPE DEVELOPMENT AND CURRENT STATUS
    # -------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("3. Prototype Development and Current Status")
    r.font.color.rgb = RGBColor(26, 54, 93)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "The project has achieved fully functional, interactive prototypes across both primary engineering and secondary high-speed telemetry interfaces. Below is the verified feature-wise implementation status matrix:"
    )

    tbl_status = doc.add_table(rows=11, cols=4)
    tbl_status.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths_status = [Inches(1.8), Inches(1.2), Inches(1.1), Inches(2.4)]
    
    headers_status = ["Feature / Subsystem", "Module & Source", "Current Status", "Verification Evidence"]
    for i, h in enumerate(headers_status):
        tbl_status.cell(0, i).paragraphs[0].text = h
    style_table_header(tbl_status.rows[0], widths_status)
    
    status_data = [
        ("Vehicle Powertrain Simulator", "simulator/vehicle_model.py", "Completed", "tests/test_simulator.py (6 passed); physical energy balance verified"),
        ("Fault Injection Engine", "simulator/fault_injector.py", "Completed", "6 fault modes (fan failure, bias, freeze, missing) verified"),
        ("Noisy Telemetry Generator", "simulator/sensor_generator.py", "Completed", "Seedable Gaussian noise and standardized JSON schema verified"),
        ("Stochastic Network Latency", "network/delay_model.py", "Completed", "Fixed, uniform, and Gaussian cellular delay profiles tested"),
        ("Bernoulli Packet Loss", "network/packet_loss.py", "Completed", "tests/test_network/test_packet_loss.py (10 passed); drop rates verified"),
        ("Age of Information Tracking", "network/aoi_tracker.py", "Completed", "Multi-channel continuous trapezoidal AoI integration verified"),
        ("Nominal Physics Twin", "digital_twin/nominal_model.py", "Completed", "State estimator and expected measurement generation verified"),
        ("Residual Anomaly Detector", "digital_twin/anomaly_detector.py", "Completed", "tests/test_digital_twin/test_anomaly_detector.py (31 passed)"),
        ("Interactive Streamlit Cockpit", "app.py (Port 8501)", "Completed", "Live controls, multi-tab layout, and campaign explorer verified"),
        ("HTML5 Telemetry Cockpit", "dashboard/server.py (Port 8050)", "Completed", "Standalone HTTP server and real-time SVG gauges verified")
    ]
    for row_idx, data in enumerate(status_data, start=1):
        row = tbl_status.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
        style_table_row(row, is_even=(row_idx % 2 == 0), col_widths=widths_status)

    doc.add_heading("3.1 Screenshot Demonstration Instructions", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "To capture live visual proof of the operational prototype for the academic evaluation panel:\n"
        "1. Launch the primary dashboard: double-click run_dashboard.bat or execute 'python -m streamlit run app.py'.\n"
        "2. Access the interface at http://localhost:8501 in any web browser.\n"
        "3. In the left sidebar under 'Thermal State on Reset', select 'Warm Operating Temp (95.0°C)'.\n"
        "4. Choose scenario 'Cooling Fan Failure' and click 'Reset Simulation'.\n"
        "5. Click 'Step (+1.0s)' five times to advance the clock to t = 5.0s.\n"
        "6. In Tab 2 ('Digital Twin Diagnostics'), capture the screen displaying the critical alert: 'CRITICAL: Cooling fan actuator failure! ECU thermostat commanded fan ON (Nom=1), but actuator relay disabled (Obs=0)' and the persistent thermal residual exceeding +5.0°C."
    )

    # -------------------------------------------------------------------------
    # REQUIREMENT 4: DATASET AND EXPERIMENTAL SETUP
    # -------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("4. Dataset and Experimental Setup")
    r.font.color.rgb = RGBColor(26, 54, 93)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "Because commercial automotive fleets treat microsecond-level raw telemetry and cellular channel traces as proprietary trade secrets, this research generates physics-grounded synthetic telemetry. The generative model couples dynamic driving cycles with differential equations and stochastic channel models."
    )

    doc.add_heading("4.1 Research Campaign Sweep Matrix (93 Runs)", level=2)
    
    tbl_exp = doc.add_table(rows=9, cols=5)
    tbl_exp.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths_exp = [Inches(1.5), Inches(0.8), Inches(1.3), Inches(1.3), Inches(1.6)]
    
    headers_exp = ["Campaign Group", "Runs", "Fault Scenario", "Packet Loss (%)", "Delay / Jitter (ms)"]
    for i, h in enumerate(headers_exp):
        tbl_exp.cell(0, i).paragraphs[0].text = h
    style_table_header(tbl_exp.rows[0], widths_exp)
    
    exp_matrix = [
        ("baseline", "3", "NORMAL", "0% (Clean Link)", "10 ms (Fixed)"),
        ("packet_loss_sweep", "15", "NORMAL", "0% to 50% (5 steps)", "10 ms (Fixed)"),
        ("delay_sweep", "15", "NORMAL", "0%", "10 to 600 ms (5 steps)"),
        ("jitter_sweep", "15", "NORMAL", "0%", "Mean 100-175 ms, Std 10-30 ms"),
        ("fault_severity_sweep", "12", "TEMP_SENSOR_BIAS", "0%", "10 ms (Bias +0°C to +30°C)"),
        ("fault_packet_loss", "15", "TEMP_SENSOR_BIAS", "0% to 50%", "10 ms (Bias +15°C)"),
        ("fault_delay", "15", "TEMP_SENSOR_BIAS", "0%", "10 to 600 ms (Bias +15°C)"),
        ("combined_impairment", "3", "TEMP_SENSOR_BIAS", "25% Loss", "80 ms Latency (Stress Test)")
    ]
    for row_idx, data in enumerate(exp_matrix, start=1):
        row = tbl_exp.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
        style_table_row(row, is_even=(row_idx % 2 == 0), col_widths=widths_exp)

    # -------------------------------------------------------------------------
    # REQUIREMENT 5: TESTING AND VALIDATION
    # -------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("5. Testing and Validation")
    r.font.color.rgb = RGBColor(26, 54, 93)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "The automated test harness comprises 171 test cases implemented across 18 test scripts using Pytest 8.4.1 and Python standard unittest. All 171 tests execute in 1.47 seconds with zero failures (100% pass rate)."
    )

    doc.add_heading("5.1 Representative Test Cases Table", level=2)
    
    tbl_tc = doc.add_table(rows=9, cols=6)
    tbl_tc.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths_tc = [Inches(1.0), Inches(1.3), Inches(1.3), Inches(1.2), Inches(1.2), Inches(0.5)]
    
    headers_tc = ["Test ID", "Target File", "Test Objective", "Input / Precondition", "Expected Outcome", "Status"]
    for i, h in enumerate(headers_tc):
        tbl_tc.cell(0, i).paragraphs[0].text = h
    style_table_header(tbl_tc.rows[0], widths_tc)
    
    tc_data = [
        ("TC-SIM-01", "vehicle_model.py", "Verify engine heat balance", "v=60 km/h, load=0.6, T0=85°C", "T in [85, 110]°C, energy conserved", "PASS"),
        ("TC-SIM-02", "vehicle_model.py", "Verify fan thermostat hysteresis", "T_eng = 96.0°C, normal mode", "Fan status = 1 (ON)", "PASS"),
        ("TC-SIM-03", "fault_injector.py", "Verify fan actuator failure", "FAN_FAILURE active at 96°C", "Fan forced to 0 (OFF)", "PASS"),
        ("TC-NET-01", "delay_model.py", "Verify transmission latency", "tau = 50 ms, frame sent at t=0s", "Delivered at t=0.05s, 0 at 0.02s", "PASS"),
        ("TC-NET-02", "packet_loss.py", "Verify Bernoulli loss rate", "p_loss = 0.20 over 500 pkts", "Empirical PDR in [70%, 90%]", "PASS"),
        ("TC-NET-03", "aoi_tracker.py", "Verify AoI sawtooth tracking", "Gen t=10.0s, arrival t=10.05s", "AoI(10.05)=0.05s, AoI(10.55)=0.55s", "PASS"),
        ("TC-DT-01", "residual_gen.py", "Verify zero-division defense", "Expected=0.0, epsilon=1e-4", "Finite relative residual, no exception", "PASS"),
        ("TC-DT-02", "anomaly_det.py", "Verify persistence counter", "+25°C persistent residual stream", "Candidate at step 1; confirmed at step 3", "PASS")
    ]
    for row_idx, data in enumerate(tc_data, start=1):
        row = tbl_tc.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
        style_table_row(row, is_even=(row_idx % 2 == 0), col_widths=widths_tc)

    # -------------------------------------------------------------------------
    # REQUIREMENT 6: QUANTITATIVE RESULTS AND PERFORMANCE ANALYSIS
    # -------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("6. Quantitative Results and Performance Analysis")
    r.font.color.rgb = RGBColor(26, 54, 93)

    doc.add_heading("6.1 Controlled 40.0-Second Comparative Fault Experiment", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "A controlled 40.0-second comparison (dt = 0.1s, 400 steps) was executed starting at warm operating temperature (T0 = 95.0°C) using digital_twin/run_controlled_experiment.py. Outputs are stored in experiments/results/controlled_fault_experiment.csv and .json:"
    )

    tbl_ctrl = doc.add_table(rows=10, cols=3)
    tbl_ctrl.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths_ctrl = [Inches(2.5), Inches(2.0), Inches(2.0)]
    
    headers_ctrl = ["Measured Metric", "Normal Operation (Healthy)", "Cooling Fan Failure (Faulted)"]
    for i, h in enumerate(headers_ctrl):
        tbl_ctrl.cell(0, i).paragraphs[0].text = h
    style_table_header(tbl_ctrl.rows[0], widths_ctrl)
    
    ctrl_data = [
        ("Initial Temperature (T0)", "95.00 °C", "95.00 °C"),
        ("Final Engine Temperature (T_40s)", "94.75 °C (Thermostat stable)", "100.40 °C (+5.65 °C excess heating)"),
        ("Mean Operating Temperature", "95.33 °C", "98.20 °C"),
        ("Maximum Operating Temperature", "95.91 °C", "100.40 °C (Monotonically escalating)"),
        ("Final Nominal Expected Temp", "94.75 °C", "94.74 °C"),
        ("Final Thermal Residual (r_temp)", "-0.005 °C", "+5.660 °C (Exceeds 5.0°C threshold)"),
        ("Cooling Fan Actuation Ratio", "100.0% (ECU commands ON)", "0.0% (Actuator stuck OFF)"),
        ("First Confirmed Anomaly Time", "None (Zero False Alarms)", "t = 0.50 s (3 persistence counts)"),
        ("Confirmed Anomaly Steps", "0 / 400 steps", "396 / 400 steps")
    ]
    for row_idx, data in enumerate(ctrl_data, start=1):
        row = tbl_ctrl.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
        style_table_row(row, is_even=(row_idx % 2 == 0), col_widths=widths_ctrl)

    doc.add_heading("6.2 Research Campaign Performance Analysis (93 Runs)", level=2)
    tbl_camp = doc.add_table(rows=7, cols=3)
    tbl_camp.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths_camp = [Inches(2.2), Inches(1.5), Inches(2.8)]
    
    headers_camp = ["Campaign Metric", "Empirical Value", "Scientific & Operational Interpretation"]
    for i, h in enumerate(headers_camp):
        tbl_camp.cell(0, i).paragraphs[0].text = h
    style_table_header(tbl_camp.rows[0], widths_camp)
    
    camp_data = [
        ("Total Telemetry Packets", "92,315 packets", "Cumulative telemetry packets across 93 20-second runs"),
        ("Delivered / Dropped Packets", "84,749 / 7,566", "Packet transmission outcomes under loss models"),
        ("Packet Delivery Ratio (PDR)", "91.80%", "Measured channel throughput across all sweep groups"),
        ("Mean Age of Information (AoI)", "0.1672 seconds", "Continuous time-average data freshness across streams"),
        ("Fault Detection Success Rate", "42 / 45 (93.3%)", "3 undetected runs were sub-threshold (<=2°C bias)"),
        ("Run-Level False Alarm Rate", "0 / 48 (0.0%)", "Zero false alerts during healthy driving due to persistence")
    ]
    for row_idx, data in enumerate(camp_data, start=1):
        row = tbl_camp.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
        style_table_row(row, is_even=(row_idx % 2 == 0), col_widths=widths_camp)

    # -------------------------------------------------------------------------
    # REQUIREMENT 7: INTEGRATION AND VERIFICATION
    # -------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("7. Integration and Verification")
    r.font.color.rgb = RGBColor(26, 54, 93)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "All subsystem integration points were verified end-to-end through automated test harnesses and live browser interaction. Below is the integration verification status table:"
    )

    tbl_int = doc.add_table(rows=7, cols=4)
    tbl_int.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths_int = [Inches(1.8), Inches(1.4), Inches(1.1), Inches(2.2)]
    
    headers_int = ["Subsystem Interface", "Data Contract / Protocol", "Status", "Verification Evidence"]
    for i, h in enumerate(headers_int):
        tbl_int.cell(0, i).paragraphs[0].text = h
    style_table_header(tbl_int.rows[0], widths_int)
    
    int_data = [
        ("Vehicle Model -> Sensor Gen", "VehicleState dataclass", "Verified", "tests/test_simulator.py; verified ground truth conversion"),
        ("Sensor Gen -> Network", "TelemetryPacket to NetworkPacket", "Verified", "tests/test_network/test_network_emulator.py; serialization verified"),
        ("Network -> Cloud State Estimator", "Time-ordered packet delivery", "Verified", "scratch/verify_implementation.py; delivery timing confirmed"),
        ("Nominal Model + Obs -> Residuals", "ExpectedMeasurement vs Obs dict", "Verified", "tests/test_digital_twin/test_residual_generator.py (8 passed)"),
        ("Residuals -> Anomaly Detector", "SensorResidual dataclass", "Verified", "tests/test_digital_twin/test_anomaly_detector.py (31 passed)"),
        ("Backend -> Streamlit UI", "DashboardSimulationSession API", "Verified", "tests/test_dashboard_manager.py (7 passed); port 8501 active")
    ]
    for row_idx, data in enumerate(int_data, start=1):
        row = tbl_int.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
        style_table_row(row, is_even=(row_idx % 2 == 0), col_widths=widths_int)

    # -------------------------------------------------------------------------
    # REQUIREMENT 8: TECHNICAL CHALLENGES AND LIMITATIONS
    # -------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("8. Technical Challenges and Limitations")
    r.font.color.rgb = RGBColor(26, 54, 93)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "During system development, four primary technical challenges were identified, debugged, and resolved:\n\n"
        "1. Fan Failure Actuator Disambiguation under Cold Start:\n"
        "   • Issue: When testing cooling fan failure from a cold engine (85.0°C), the ECU thermostat naturally commands the fan OFF. Forcing fan failure caused no immediate physical discrepancy until the engine heated past 95°C.\n"
        "   • Resolution: Implemented a 'Thermal State on Reset' feature allowing tests to initialize at warm operating temperature (95.0°C) with ECU thermostat active, coupled with categorical actuator checking (F_hat != F_obs). This provides instant fault detection at t = 0.50s.\n\n"
        "2. Multi-Sensor Age of Information Tracking:\n"
        "   • Issue: Early versions tracked AoI solely on the temperature stream, leaving vehicle speed, RPM, and engine load unmonitored during cellular dropouts.\n"
        "   • Resolution: Extended AoITracker to maintain individual SensorAoIState tracking instances across all five supported channels.\n\n"
        "3. Transient Network Jitter Inducing False Alarms:\n"
        "   • Issue: Intermittent delivery bursts caused numerical derivative spikes that tripped static thresholds.\n"
        "   • Resolution: Implemented a two-level threshold hierarchy (Warning: 2.0°C, Anomaly: 5.0°C) with a persistence requirement (N = 3 consecutive cycles), achieving a 0.0% false alarm rate across 48 healthy campaign runs.\n\n"
        "4. Current Limitations:\n"
        "   • Telemetry transmission remains passive and fixed-rate.\n"
        "   • The baseline detector identifies that an anomaly exists but cannot compute a multi-hypothesis probability distribution over competing causes (e.g. fan failure vs coolant leak)."
    )

    # -------------------------------------------------------------------------
    # REQUIREMENT 9: MILESTONE ACHIEVEMENT AND FUTURE WORK
    # -------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("9. Milestone Achievement and Future Work")
    r.font.color.rgb = RGBColor(26, 54, 93)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "The project has satisfied 100% of Phase III milestones. Below is the milestone matrix and the prioritized research roadmap for Phase IV (Active Digital Twin Core):"
    )

    tbl_ms = doc.add_table(rows=7, cols=4)
    tbl_ms.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths_ms = [Inches(1.5), Inches(1.5), Inches(1.0), Inches(2.5)]
    
    headers_ms = ["Milestone", "Target Deliverable", "Status", "Completion Evidence"]
    for i, h in enumerate(headers_ms):
        tbl_ms.cell(0, i).paragraphs[0].text = h
    style_table_header(tbl_ms.rows[0], widths_ms)
    
    ms_data = [
        ("Milestone 1", "Vehicle Powertrain Simulator", "Completed", "simulator/ module; 171 tests passing"),
        ("Milestone 2", "Network Impairment & AoI", "Completed", "network/ module; delay, loss, AoI verified"),
        ("Milestone 3", "Cloud Digital Twin Backend", "Completed", "digital_twin/ nominal model & residual generator"),
        ("Milestone 4", "Experimental Benchmark Suite", "Completed", "93-run campaign dataset in experiments/results/"),
        ("Milestone 5", "Interactive Cockpit Dashboards", "Completed", "Streamlit (:8501) and HTML5 Cockpit (:8050) active"),
        ("Milestone 6", "Bayesian Active Digital Twin", "Planned", "Phase IV Roadmap: belief tracking & VoI action engine")
    ]
    for row_idx, data in enumerate(ms_data, start=1):
        row = tbl_ms.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
        style_table_row(row, is_even=(row_idx % 2 == 0), col_widths=widths_ms)

    doc.add_heading("9.1 Phase IV Research Roadmap (Active Digital Twin Enhancements)", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "To achieve the ultimate research vision ('Uncertainty- and Cost-Aware Active Digital Twin'), Phase IV will introduce:\n\n"
        "1. Multi-Hypothesis Bayesian Fault-Belief Tracking:\n"
        "   Maintain a probability belief vector b_k over competing hypotheses:\n"
        "   H0: Healthy under dynamic load | H1: Cooling fan failure | H2: Temp sensor bias | H3: Coolant fluid leak\n"
        "   b_k(Hi) = P(Hi | y_1:k) = [ P(y_k | Hi, y_1:k-1) * b_k-1(Hi) ] / [ sum_j P(y_k | Hj, y_1:k-1) * b_k-1(Hj) ]\n\n"
        "2. Diagnostic Uncertainty Quantification via Belief Entropy:\n"
        "   H(b_k) = - sum_i=0^3 b_k(Hi) * log2 b_k(Hi)\n"
        "   Quantifies diagnostic confusion in continuous bits to trigger active probing.\n\n"
        "3. Cost-Aware Active Diagnostic Decision Engine:\n"
        "   Select action a_k* in {PASSIVE_WAIT, BURST_TELEMETRY, CYCLE_FAN_ACTUATOR, DECLARE_FAULT} maximizing:\n"
        "   a_k* = argmax_a [ E[Delta H(b_k, a)] - lambda_cost * C_comm(a) - lambda_aoi * AoI_k ]"
    )

    # -------------------------------------------------------------------------
    # CONCLUSION & SUBMISSION CHECKLIST
    # -------------------------------------------------------------------------
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("10. Conclusion and Submission Readiness Checklist")
    r.font.color.rgb = RGBColor(26, 54, 93)

    chk_text = (
        "[X] Complete source code organized, clean, and synchronized on Git main (Commit: 46996fa)\n"
        "[X] Zero unpushed commits or exposed credentials (.env protected in .gitignore)\n"
        "[X] 171 automated unit and integration tests passing with 100% success rate in 1.47 seconds\n"
        "[X] Reproducible 40-second controlled cooling fan fault comparison documented (CSV & JSON)\n"
        "[X] 93-replicate research campaign across eight impairment sweeps analyzed and validated\n"
        "[X] Both Streamlit (:8501) and HTML5 Cockpit (:8050) operational and launchable via batch scripts\n"
        "[X] Complete Capstone_Project_Comprehensive_Guide.pdf generated via ReportLab\n"
        "[X] All nine college evaluation criteria satisfied with empirical evidence and source references"
    )
    add_callout(doc, chk_text, title="CAPSTONE PROJECT SUBMISSION READINESS VERIFIED", border_color="2B6CB0", bg_color="F7FAFC")

    # Save document
    doc.save(output_path)
    print(f"Successfully generated Word report: {output_path} ({os.path.getsize(output_path)} bytes)")

if __name__ == "__main__":
    out_file = os.path.join(r"C:\Users\ANNU TIWARI\Desktop\Capstone project", "Capstone_Project_Phase_III_Midterm_Report.docx")
    build_word_report(out_file)
